#!/usr/bin/env python3
"""把 docs/ 下的 Markdown 交付文档转成 .docx（标题/列表/表格/加粗）。

用法：python3 scripts/md2docx.py docs/xxx.md docs/xxx.docx
"""

import re
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt

BOLD_RE = re.compile(r"\*\*(.+?)\*\*")
CODE_RE = re.compile(r"`([^`]+)`")


def set_base_font(document):
    style = document.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")


def add_runs(paragraph, text):
    """支持 **加粗** 与 `代码` 两种行内标记。"""
    text = CODE_RE.sub(r"\1", text)
    position = 0
    for match in BOLD_RE.finditer(text):
        if match.start() > position:
            paragraph.add_run(text[position:match.start()])
        run = paragraph.add_run(match.group(1))
        run.bold = True
        position = match.end()
    if position < len(text):
        paragraph.add_run(text[position:])


def is_table_separator(line):
    return bool(re.fullmatch(r"\|[\s:\-|]+\|", line.strip()))


def split_row(line):
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def convert(md_path, docx_path):
    document = Document()
    set_base_font(document)
    lines = open(md_path, encoding="utf-8").read().splitlines()
    index = 0
    while index < len(lines):
        stripped = lines[index].strip()
        if not stripped:
            index += 1
            continue
        if stripped.startswith("```"):
            index += 1
            while index < len(lines) and not lines[index].strip().startswith("```"):
                document.add_paragraph(lines[index])
                index += 1
            index += 1
            continue
        if stripped.startswith("|") and index + 1 < len(lines) and is_table_separator(lines[index + 1]):
            header = split_row(stripped)
            index += 2
            rows = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                rows.append(split_row(lines[index]))
                index += 1
            table = document.add_table(rows=1, cols=len(header))
            table.style = "Table Grid"
            for column, text in enumerate(header):
                cell = table.rows[0].cells[column]
                cell.text = ""
                add_runs(cell.paragraphs[0], text)
                for run in cell.paragraphs[0].runs:
                    run.bold = True
            for row in rows:
                cells = table.add_row().cells
                for column, text in enumerate(row[:len(header)]):
                    cells[column].text = ""
                    add_runs(cells[column].paragraphs[0], text)
            continue
        heading = re.match(r"^(#{1,4})\s+(.*)$", stripped)
        if heading:
            paragraph = document.add_heading(level=len(heading.group(1)))
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            add_runs(paragraph, heading.group(2))
            index += 1
            continue
        if re.match(r"^[-*]\s+", stripped):
            paragraph = document.add_paragraph(style="List Bullet")
            add_runs(paragraph, re.sub(r"^[-*]\s+", "", stripped))
            index += 1
            continue
        if re.match(r"^\d+\.\s+", stripped):
            paragraph = document.add_paragraph(style="List Number")
            add_runs(paragraph, re.sub(r"^\d+\.\s+", "", stripped))
            index += 1
            continue
        if stripped.startswith(">"):
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Pt(18)
            add_runs(paragraph, stripped.lstrip("> ").strip())
            index += 1
            continue
        paragraph = document.add_paragraph()
        add_runs(paragraph, stripped)
        index += 1
    document.save(docx_path)


if __name__ == "__main__":
    if len(sys.argv) != 3:
        raise SystemExit("用法：python3 scripts/md2docx.py <input.md> <output.docx>")
    convert(sys.argv[1], sys.argv[2])
    print(f"已生成 {sys.argv[2]}")
